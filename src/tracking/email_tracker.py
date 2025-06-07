import pandas as pd
import datetime
from typing import List, Dict, Optional
import os
from src.utils.logger import logger


class EmailTracker:
    """Tracks email sending activities and exports to Excel."""
    
    def __init__(self, output_file: str = "email_tracking.xlsx"):
        """
        Initialize the email tracker.
        
        Args:
            output_file: Path to the output Excel file
        """
        self.output_file = output_file
        self.tracking_data = []
        
        # Load existing data if file exists
        if os.path.exists(output_file):
            try:
                df = pd.read_excel(output_file, sheet_name='Email Tracking')
                self.tracking_data = df.to_dict('records')
                logger.info(f"Loaded {len(self.tracking_data)} existing records from {output_file}")
            except Exception as e:
                logger.warning(f"Could not load existing data from {output_file}: {str(e)}")
        
    def extract_name_from_email(self, email: str) -> Dict[str, str]:
        """
        Extract first and last name from email address.
        
        Args:
            email: Email address
            
        Returns:
            Dict with first_name and last_name
        """
        try:
            # Extract the part before @
            username = email.split("@")[0]
            
            # Replace common separators with spaces
            name_part = username.replace(".", " ").replace("_", " ").replace("-", " ")
            
            # Split into parts and capitalize
            name_parts = [part.capitalize() for part in name_part.split() if part]
            
            if len(name_parts) == 0:
                return {"first_name": "Unknown", "last_name": ""}
            elif len(name_parts) == 1:
                return {"first_name": name_parts[0], "last_name": ""}
            else:
                return {"first_name": name_parts[0], "last_name": " ".join(name_parts[1:])}
                
        except Exception as e:
            logger.warning(f"Error extracting name from {email}: {str(e)}")
            return {"first_name": "Unknown", "last_name": ""}
    
    def log_email_attempt(self, email: str, subject: str, status: str, error_message: str = None, message_id: str = None, is_followup: bool = False) -> None:
        """
        Log an email sending attempt.
        
        Args:
            email: Recipient email address
            subject: Email subject
            status: 'success' or 'failed'
            error_message: Error message if failed
            message_id: Gmail message ID for threading
            is_followup: Whether this is a follow-up email
        """
        try:
            name_info = self.extract_name_from_email(email)
            
            tracking_entry = {
                "timestamp": datetime.datetime.now(),
                "email": email,
                "first_name": name_info["first_name"],
                "last_name": name_info["last_name"],
                "full_name": f"{name_info['first_name']} {name_info['last_name']}".strip(),
                "subject": subject,
                "status": status,
                "error_message": error_message or "",
                "domain": email.split("@")[1] if "@" in email else "",
                "date": datetime.datetime.now().strftime("%Y-%m-%d"),
                "time": datetime.datetime.now().strftime("%H:%M:%S"),
                "message_id": message_id or "",
                "is_followup": is_followup,
                "followup_sent": False if not is_followup else True
            }
            
            self.tracking_data.append(tracking_entry)
            
        except Exception as e:
            logger.error(f"Error logging email attempt: {str(e)}")
    
    def export_to_excel(self) -> None:
        """Export tracking data to Excel with formatting."""
        try:
            if not self.tracking_data:
                logger.warning("No tracking data to export")
                return
            
            # Create DataFrame and sort by timestamp
            df = pd.DataFrame(self.tracking_data)
            df = df.sort_values('timestamp', ascending=False)  # Sort by timestamp, newest first
            
            # Remove any duplicate entries based on email and timestamp
            df = df.drop_duplicates(subset=['email', 'timestamp'], keep='first')
            
            # Create Excel writer with formatting
            with pd.ExcelWriter(self.output_file, engine='openpyxl') as writer:
                # Write main data
                df.to_excel(writer, sheet_name='Email Tracking', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Email Tracking']
                
                # Format headers
                header_font = openpyxl.styles.Font(bold=True, color='FFFFFF')
                header_fill = openpyxl.styles.PatternFill(start_color='366092', end_color='366092', fill_type='solid')
                
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add conditional formatting for status
                from openpyxl.styles import PatternFill
                from openpyxl.formatting.rule import CellIsRule
                
                success_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                fail_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                
                status_column = None
                for idx, col in enumerate(df.columns, 1):
                    if col == 'status':
                        status_column = openpyxl.utils.get_column_letter(idx)
                        break
                
                if status_column:
                    worksheet.conditional_formatting.add(f'{status_column}2:{status_column}{len(df)+1}',
                                                       CellIsRule(operator='equal', formula=['"success"'], fill=success_fill))
                    worksheet.conditional_formatting.add(f'{status_column}2:{status_column}{len(df)+1}',
                                                       CellIsRule(operator='equal', formula=['"failed"'], fill=fail_fill))
                
                # Create summary sheet
                summary_data = self._create_summary_data(df)
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
                # Format summary sheet
                summary_worksheet = writer.sheets['Summary']
                for cell in summary_worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                
                for column in summary_worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 30)
                    summary_worksheet.column_dimensions[column_letter].width = adjusted_width
            
            logger.info(f"Email tracking data exported to {self.output_file}")
            
        except Exception as e:
            logger.error(f"Error exporting to Excel: {str(e)}")
            # Fallback to CSV
            try:
                df = pd.DataFrame(self.tracking_data)
                csv_file = self.output_file.replace('.xlsx', '.csv')
                df.to_csv(csv_file, index=False)
                logger.info(f"Exported to CSV instead: {csv_file}")
            except Exception as csv_error:
                logger.error(f"Failed to export to CSV as well: {str(csv_error)}")
    
    def _create_summary_data(self, df: pd.DataFrame) -> List[Dict]:
        """Create summary statistics."""
        try:
            total_emails = len(df)
            successful = len(df[df['status'] == 'success'])
            failed = len(df[df['status'] == 'failed'])
            success_rate = (successful / total_emails * 100) if total_emails > 0 else 0
            
            # Domain breakdown
            domain_counts = df['domain'].value_counts().head(10)
            
            summary_data = [
                {"Metric": "Total Emails Sent", "Value": total_emails},
                {"Metric": "Successful", "Value": successful},
                {"Metric": "Failed", "Value": failed},
                {"Metric": "Success Rate (%)", "Value": f"{success_rate:.1f}%"},
                {"Metric": "Start Time", "Value": df['timestamp'].min().strftime("%Y-%m-%d %H:%M:%S") if not df.empty else ""},
                {"Metric": "End Time", "Value": df['timestamp'].max().strftime("%Y-%m-%d %H:%M:%S") if not df.empty else ""},
                {"Metric": "", "Value": ""},  # Empty row
                {"Metric": "Top Domains", "Value": ""},
            ]
            
            for domain, count in domain_counts.items():
                summary_data.append({"Metric": f"  {domain}", "Value": count})
            
            return summary_data
            
        except Exception as e:
            logger.error(f"Error creating summary data: {str(e)}")
            return [{"Metric": "Error creating summary", "Value": str(e)}]
    
    def get_statistics(self) -> Dict:
        """Get current statistics."""
        try:
            if not self.tracking_data:
                return {"total": 0, "successful": 0, "failed": 0, "success_rate": 0}
            
            total = len(self.tracking_data)
            successful = sum(1 for entry in self.tracking_data if entry["status"] == "success")
            failed = total - successful
            success_rate = (successful / total * 100) if total > 0 else 0
            
            return {
                "total": total,
                "successful": successful,
                "failed": failed,
                "success_rate": success_rate
            }
            
        except Exception as e:
            logger.error(f"Error getting statistics: {str(e)}")
            return {"total": 0, "successful": 0, "failed": 0, "success_rate": 0}
    
    def clear_data(self) -> None:
        """Clear all tracking data."""
        self.tracking_data = []
        logger.info("Tracking data cleared")
    
    def mark_followup_sent(self, email: str, original_timestamp: datetime.datetime) -> None:
        """Mark that a follow-up has been sent for a specific email record."""
        try:
            for entry in self.tracking_data:
                if (entry["email"] == email and 
                    entry["timestamp"] == original_timestamp and 
                    not entry["is_followup"]):
                    entry["followup_sent"] = True
                    logger.info(f"Marked follow-up as sent for {email}")
                    break
        except Exception as e:
            logger.error(f"Error marking follow-up as sent: {str(e)}")
    
    def get_emails_needing_followup(self, days_threshold: int = 2) -> List[Dict]:
        """
        Get emails that need follow-up (sent X days ago, successful, no follow-up sent yet).
        
        Args:
            days_threshold: Number of days to wait before sending follow-up
            
        Returns:
            List of email records that need follow-up
        """
        try:
            cutoff_date = datetime.datetime.now() - datetime.timedelta(days=days_threshold)
            followup_needed = []
            
            for entry in self.tracking_data:
                if (entry["status"] == "success" and 
                    not entry["is_followup"] and 
                    not entry["followup_sent"] and
                    entry["timestamp"] < cutoff_date):
                    followup_needed.append(entry)
            
            return followup_needed
            
        except Exception as e:
            logger.error(f"Error getting emails needing follow-up: {str(e)}")
            return []


# Add openpyxl import for Excel formatting
try:
    import openpyxl
    import openpyxl.styles
    import openpyxl.utils
    import openpyxl.formatting.rule
except ImportError:
    logger.warning("openpyxl not installed. Excel formatting will be limited.") 